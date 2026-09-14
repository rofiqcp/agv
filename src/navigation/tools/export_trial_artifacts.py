#!/usr/bin/env python3
"""Export one experiment trial to report-ready XLSX and Matplotlib PNG files."""
from __future__ import annotations

import argparse
import csv
import json
import math
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


def load_json(path: str, default):
    p = Path(path)
    if not p.is_file():
        return default
    return json.loads(p.read_text(encoding="utf-8"))


def to_float(value):
    try:
        x = float(value)
        return x if math.isfinite(x) else None
    except (TypeError, ValueError):
        return None

def read_rows(path: str):
    with open(path, newline="", encoding="utf-8") as handle:
        reader = csv.DictReader(handle)
        return list(reader), list(reader.fieldnames or [])


def excel_value(value):
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    return value


def style_sheet(ws):
    fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"
    for idx, col in enumerate(ws.columns, 1):
        width = min(42, max(10, max(len(str(c.value or "")) for c in col) + 2))
        ws.column_dimensions[get_column_letter(idx)].width = width


def write_workbook(rows, columns, summary, trials, output, table_csv_paths=None, spec=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Raw Data"
    ws.append(columns)
    for row in rows:
        ws.append([row.get(c, "") for c in columns])
    style_sheet(ws)

    sm = wb.create_sheet("Trial Summary")
    sm.append(["Parameter", "Value"])
    for key, value in summary.items():
        sm.append([key, json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value])
    style_sheet(sm)

    recap = wb.create_sheet("Trials Recap")
    trial_rows = trials if isinstance(trials, list) else []
    keys = []
    for trial in trial_rows:
        for key in trial.keys():
            if key not in keys:
                keys.append(key)
    if keys:
        recap.append(keys)
        for trial in trial_rows:
            recap.append([excel_value(trial.get(k, "")) for k in keys])
    else:
        recap.append(["No saved trials"])
    style_sheet(recap)

    for idx, table_path in enumerate(table_csv_paths or [], 1):
        try:
            table_rows, table_columns = read_rows(table_path)
        except (OSError, csv.Error, UnicodeDecodeError):
            continue
        tws = wb.create_sheet(f"GUI Table {idx}")
        if table_columns:
            tws.append(table_columns)
            for row in table_rows:
                tws.append([row.get(c, "") for c in table_columns])
        else:
            tws.append(["No table rows"])
        style_sheet(tws)

    # Extra analysis tables are export-only. They do not alter the main GUI tables.
    spec = spec or {}
    minute_rows = build_per_minute_summary(rows, spec)
    mws = wb.create_sheet("Per-Minute Mean")
    mws.append(["Graph", "Window", "Series", "Mean", "Std", "N", "Source", "Method"])
    for r in minute_rows:
        mws.append([r[k] for k in ("graph", "window", "series", "mean", "std", "n", "source", "method")])
    if not minute_rows:
        mws.append(["No numeric time-series samples"])
    style_sheet(mws)

    tws = wb.create_sheet("Trial Mean")
    tws.append(["Series", "Mean", "Std", "Min", "Max", "P95", "N", "Scope"])
    trial_stats = summary.get("series_stats", {}) if isinstance(summary, dict) else {}
    if isinstance(trial_stats, dict) and trial_stats:
        for label, st in trial_stats.items():
            if not isinstance(st, dict):
                continue
            tws.append([label, st.get("mean", ""), st.get("std", ""), st.get("min", ""), st.get("max", ""), st.get("p95", ""), st.get("n", ""), "whole trial"])
    else:
        tws.append(["No client trial statistics"])
    style_sheet(tws)

    rws = wb.create_sheet("Repeat Mean")
    rws.append(["Series", "Repeat Mean", "Repeat Std", "N Trials", "Scope"])
    repeat = build_repeat_summary(trials)
    for label, st in repeat.items():
        rws.append([label, st["mean"], st["std"], st["n"], "same experiment/task"])
    if not repeat:
        rws.append(["Need >=2 compatible trials with series_means"])
    style_sheet(rws)
    wb.save(output)


def series_xy(rows, x_key, y_key):
    xs, ys = [], []
    for row in rows:
        x, y = to_float(row.get(x_key)), to_float(row.get(y_key))
        if x is None or y is None:
            continue
        xs.append(x); ys.append(y)
    return xs, ys


def _json_points(value):
    if not value:
        return []
    try:
        data = json.loads(value) if isinstance(value, str) else value
    except (TypeError, json.JSONDecodeError):
        return []
    out = []
    for p in data if isinstance(data, list) else []:
        if isinstance(p, (list, tuple)) and len(p) >= 2:
            x, y = to_float(p[0]), to_float(p[1])
            if x is not None and y is not None:
                out.append((x, y))
    return out


def _localize_lon_lat(xs, ys):
    if not xs or not ys:
        return xs, ys
    lon0, lat0 = xs[0], ys[0]
    radius = 6378137.0
    lat0_rad = math.radians(lat0)
    east = [(lon - lon0) * math.pi / 180.0 * radius * math.cos(lat0_rad) for lon in xs]
    north = [(lat - lat0) * math.pi / 180.0 * radius for lat in ys]
    return east, north


def _mean_std(values):
    vals = [v for v in values if v is not None and math.isfinite(v)]
    if not vals:
        return None
    mean = sum(vals) / len(vals)
    var = sum((v - mean) ** 2 for v in vals) / len(vals)
    ordered = sorted(vals)
    p95 = ordered[min(len(ordered) - 1, int((len(ordered) - 1) * 0.95))]
    return {"mean": mean, "std": math.sqrt(var), "min": ordered[0], "max": ordered[-1], "p95": p95, "n": len(vals)}


def build_per_minute_summary(rows, spec):
    live = spec.get("live_series", {}) if isinstance(spec, dict) else {}
    out = []
    for gi, graph in enumerate(spec.get("graphs", []) if isinstance(spec, dict) else [], 1):
        if graph.get("type", "time_series") != "time_series":
            continue
        title = graph.get("title") or f"Graph {gi}"
        for label in graph.get("series", []):
            key = live.get(label, label)
            buckets = {}
            for row in rows:
                t, y = to_float(row.get("elapsed_s")), to_float(row.get(key))
                if t is None or y is None:
                    continue
                minute = max(0, int(t // 60.0))
                buckets.setdefault(minute, []).append(y)
            for minute in sorted(buckets):
                st = _mean_std(buckets[minute])
                if not st:
                    continue
                out.append({"graph": title, "window": f"{minute*60}-{(minute+1)*60} s", "series": label,
                            "mean": st["mean"], "std": st["std"], "n": st["n"], "source": key,
                            "method": "arithmetic mean per fixed 60 s window"})
    return out


def build_repeat_summary(trials):
    grouped = {}
    for trial in trials if isinstance(trials, list) else []:
        means = trial.get("series_means", {}) if isinstance(trial, dict) else {}
        if not isinstance(means, dict):
            continue
        for label, value in means.items():
            v = to_float(value)
            if v is not None:
                grouped.setdefault(label, []).append(v)
    out = {}
    for label, values in grouped.items():
        if len(values) < 2:
            continue
        st = _mean_std(values)
        if st:
            out[label] = st
    return out


def save_summary_graphs(rows, spec, trials, output_prefix):
    live = spec.get("live_series", {}) if isinstance(spec, dict) else {}
    paths, titles = [], []
    repeat = build_repeat_summary(trials)
    for gi, graph in enumerate(spec.get("graphs", []) if isinstance(spec, dict) else [], 1):
        if graph.get("type", "time_series") != "time_series":
            continue
        labels = list(graph.get("series", []))
        if not labels:
            continue
        base_title = graph.get("title") or f"Graph {gi}"

        # Per-minute mean: 60-second fixed windows.
        fig, ax = plt.subplots(figsize=(8.0, 4.8), facecolor="white")
        plotted = False
        for label in labels:
            key = live.get(label, label)
            buckets = {}
            for row in rows:
                t, y = to_float(row.get("elapsed_s")), to_float(row.get(key))
                if t is None or y is None:
                    continue
                m = max(0, int(t // 60.0)); buckets.setdefault(m, []).append(y)
            xs, ys = [], []
            for m in sorted(buckets):
                st = _mean_std(buckets[m])
                if st:
                    xs.append(m + 1); ys.append(st["mean"])
            if xs:
                ax.plot(xs, ys, marker="o", linewidth=1.5, label=label); plotted = True
        if plotted:
            title = f"{base_title} — Per-minute mean (60 s window)"
            ax.set_title(title); ax.set_xlabel("Minute window [-]"); ax.set_ylabel(graph.get("yLabel") or "Mean value")
            ax.grid(True, alpha=0.25); ax.legend(loc="best"); fig.tight_layout()
            path = f"{output_prefix}_G{gi:02d}_PER_MINUTE.png"; fig.savefig(path, dpi=180, facecolor="white", bbox_inches="tight")
            paths.append(path); titles.append(title)
        plt.close(fig)

        # Trial mean: mean over the complete current run.
        trial_labels, trial_means = [], []
        for label in labels:
            key = live.get(label, label)
            st = _mean_std([to_float(r.get(key)) for r in rows])
            if st:
                trial_labels.append(label); trial_means.append(st["mean"])
        if trial_means:
            fig, ax = plt.subplots(figsize=(8.0, 4.8), facecolor="white")
            ax.bar(range(len(trial_means)), trial_means); ax.set_xticks(range(len(trial_labels)), trial_labels, rotation=20, ha="right")
            title = f"{base_title} — Trial mean (whole run)"
            ax.set_title(title); ax.set_ylabel(graph.get("yLabel") or "Mean value"); ax.grid(True, axis="y", alpha=0.25); fig.tight_layout()
            path = f"{output_prefix}_G{gi:02d}_TRIAL_MEAN.png"; fig.savefig(path, dpi=180, facecolor="white", bbox_inches="tight"); plt.close(fig)
            paths.append(path); titles.append(title)

        # Repeat mean: aggregate trial means from the same experiment/task.
        rep_labels, rep_means, rep_std, rep_n = [], [], [], []
        for label in labels:
            st = repeat.get(label)
            if not st:
                continue
            rep_labels.append(label); rep_means.append(st["mean"]); rep_std.append(st["std"]); rep_n.append(st["n"])
        if rep_means:
            fig, ax = plt.subplots(figsize=(8.0, 4.8), facecolor="white")
            ax.bar(range(len(rep_means)), rep_means, yerr=rep_std, capsize=4); ax.set_xticks(range(len(rep_labels)), rep_labels, rotation=20, ha="right")
            nmin, nmax = min(rep_n), max(rep_n)
            scope = f"{nmin} repeat" if nmin == nmax else f"{nmin}-{nmax} repeat"
            title = f"{base_title} — Repeat mean ({scope})"
            ax.set_title(title); ax.set_ylabel(graph.get("yLabel") or "Mean of trial means"); ax.grid(True, axis="y", alpha=0.25); fig.tight_layout()
            path = f"{output_prefix}_G{gi:02d}_REPEAT_MEAN.png"; fig.savefig(path, dpi=180, facecolor="white", bbox_inches="tight"); plt.close(fig)
            paths.append(path); titles.append(title)
    return paths, titles


def save_graphs(rows, spec, output_prefix):
    live = spec.get("live_series", {})
    graphs = spec.get("graphs", [])
    paths = []
    for index, graph in enumerate(graphs, 1):
        fig, ax = plt.subplots(figsize=(8.0, 4.8), facecolor="white")
        ax.set_facecolor("white")
        kind = graph.get("type", "time_series")
        plotted = False
        if kind == "scatter":
            x_key = graph.get("xSeries", "")
            y_key = graph.get("ySeries", "")
            xs, ys = series_xy(rows, x_key, y_key)
            if x_key == "gnss_fix.lon" and y_key == "gnss_fix.lat":
                xs, ys = _localize_lon_lat(xs, ys)
            if xs:
                ax.plot(xs, ys, marker=".", linewidth=1.3, label="Samples")
                plotted = True
        elif kind == "multi_scatter":
            for pair in graph.get("pairs", []):
                if not isinstance(pair, list) or len(pair) < 2:
                    continue
                x_label, y_label = pair[0], pair[1]
                x_key, y_key = live.get(x_label, x_label), live.get(y_label, y_label)
                xs, ys = series_xy(rows, x_key, y_key)
                if xs:
                    ax.plot(xs, ys, linewidth=1.4, label=f"{x_label} / {y_label}")
                    plotted = True
        elif kind == "path":
            path_key = graph.get("pathKey") or "nav_path"
            points = []
            for row in reversed(rows):
                points = _json_points(row.get(path_key + ".points"))
                if points:
                    break
            if points:
                ax.plot([p[0] for p in points], [p[1] for p in points], linewidth=1.8, label=path_key)
                plotted = True
                ax.set_aspect("equal", adjustable="datalim")
        else:
            elapsed = [to_float(r.get("elapsed_s")) for r in rows]
            for label in graph.get("series", []):
                key = live.get(label, label)
                xs, ys = [], []
                for t, row in zip(elapsed, rows):
                    y = to_float(row.get(key))
                    if t is None or y is None:
                        continue
                    xs.append(t); ys.append(y)
                if xs:
                    ax.plot(xs, ys, linewidth=1.6, label=label)
                    plotted = True
        ax.set_xlabel(graph.get("xLabel") or ("Time [s]" if kind not in ("scatter", "multi_scatter", "path") else "X"))
        ax.set_ylabel(graph.get("yLabel") or "Value")
        ax.grid(True, alpha=0.25)
        if plotted:
            ax.legend(loc="best")
        else:
            ax.text(0.5, 0.5, "No numeric samples in server raw CSV\n(browser PNG may override this fallback)",
                    ha="center", va="center", transform=ax.transAxes, fontsize=9)
        title = graph.get("title") or f"Trial Graph {index}"
        ax.set_title(title)
        fig.tight_layout()
        path = f"{output_prefix}_G{index:02d}.png"
        fig.savefig(path, dpi=180, facecolor="white", bbox_inches="tight")
        plt.close(fig)
        paths.append(path)
    return paths

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--csv", required=True)
    ap.add_argument("--summary", required=True)
    ap.add_argument("--trials", required=True)
    ap.add_argument("--spec", required=True)
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--png-prefix", required=True)
    ap.add_argument("--table-csv", action="append", default=[])
    args = ap.parse_args()

    rows, columns = read_rows(args.csv)
    summary = load_json(args.summary, {})
    trials = load_json(args.trials, [])
    spec = load_json(args.spec, {})
    write_workbook(rows, columns, summary, trials, args.xlsx, args.table_csv, spec)
    pngs = save_graphs(rows, spec, args.png_prefix)
    base_titles = [(g.get("title") or f"Trial Graph {i}") for i, g in enumerate(spec.get("graphs", []), 1)]
    extra_pngs, extra_titles = save_summary_graphs(rows, spec, trials, args.png_prefix)
    pngs.extend(extra_pngs)
    print(json.dumps({"xlsx": args.xlsx, "pngs": pngs, "png_titles": base_titles + extra_titles}, ensure_ascii=False))


if __name__ == "__main__":
    main()
