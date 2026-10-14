#!/usr/bin/env python3
from __future__ import annotations
import argparse,csv,json,math,re,statistics,zipfile,shutil
from pathlib import Path
from datetime import datetime,timezone,timedelta
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

TZ=timezone(timedelta(hours=7))

TABLES={
"4.2.1":("Frequency_EKF_Lokal",["Detik [s]","Variasi","Frequency","RMSE Residual vx ESC","RMSE Residual vx GNSS","Mean EKF age","P95 EKF age","EKF continuity [%]"]),
"4.2.2":("Sensor_Timeout_EKF_Lokal",["Detik [s]","Variasi","Sensor timeout","RMSE Residual vx ESC","RMSE Residual vx GNSS","Mean EKF age","P95 EKF age","EKF continuity [%]"]),
"4.2.3":("Measurement_Covariance_EKF_Lokal",["Detik [s]","Variasi","R ESC vx","R GNSS vx","R IMU yaw","R ESC wz","R IMU wz","RMSE Residual vx ESC","RMSE Residual vx GNSS"]),
"4.2.4":("Process_Noise_EKF_Lokal",["Detik [s]","Variasi","Qx","Qy","Qyaw","Qvx","Qwz","Mean sigma X","Mean sigma Y","RMSE Residual vx ESC","RMSE Residual vx GNSS"]),
"4.2.5":("Predict_Current_Time_EKF_Lokal",["Detik [s]","Variasi","Predict to current time","RMSE Residual vx ESC","RMSE Residual vx GNSS","Mean EKF age","P95 EKF age","EKF continuity [%]"]),
"4.3.1":("Frequency_EKF_Global",["Detik [s]","Variasi","Frequency","RMSE Residual XY GNSS [m]","Mean EKF age","P95 EKF age","EKF global continuity [%]","GNSS availability [%]"]),
"4.3.2":("Sensor_Timeout_EKF_Global",["Detik [s]","Variasi","Sensor timeout","GNSS availability [%]","Longest GNSS gap [s]","EKF global continuity [%]","RMSE Residual XY GNSS [m]","Mean EKF age"]),
"4.3.3":("Measurement_Covariance_GNSS",["Detik [s]","Variasi","Mean hAcc","R GNSS X","R GNSS Y","R GNSS vx","R COG yaw","RMSE Residual XY GNSS [m]","Max position jump [m]"]),
"4.3.4":("Process_Noise_EKF_Global",["Detik [s]","Variasi","Qx","Qy","Mean sigma X","Mean sigma Y","RMSE Residual XY GNSS [m]","Max position jump [m]"]),
"4.3.5":("Predict_Current_Time_EKF_Global",["Detik [s]","Variasi","Predict to current time","RMSE Residual XY GNSS [m]","Max position jump [m]","Mean EKF age","P95 EKF age","EKF global continuity [%]"]),
"4.4.1":("Minimum_Turning_Radius",["Detik [s]","Variasi","minimum_turning_radius candidate [m]","Planning time","Path length","Path endpoint error","Minimum clearance [m]","Planner success"]),
"4.4.2":("Downsampling",["Detik [s]","Variasi","Downsample costmap","Downsampling factor","Planning time","Path length","Minimum clearance [m]","Planner success"]),
"4.4.3":("Planner_Tolerance",["Detik [s]","Variasi","Planner tolerance (m)","Path endpoint error","Planning time","Planner success","Path length"]),
"4.4.4":("Cost_Penalty",["Detik [s]","Variasi","Cost penalty","Minimum clearance [m]","Path length","Planning time","Planner success"]),
"4.4.5":("Goal_Checker",["Detik [s]","Variasi","XY goal tolerance (m)","Yaw goal tolerance (rad)","Stopped linear velocity (m/s)","Stopped yaw rate (rad/s)","Final distance to goal","Minimum distance to goal","Final speed abs","Final yaw rate abs","Measured goal pass","Nav2 result","Time-to-goal"]),
}

def f(row,key,default=None):
    try:
        v=float(row.get(key,""))
        return v if math.isfinite(v) else default
    except:return default

def b(row,key,default=False):
    v=str(row.get(key,"")).strip().lower()
    if v in {"true","1","yes"}: return True
    if v in {"false","0","no"}: return False
    return default

def p95(a):
    if not a:return None
    s=sorted(a);i=max(0,min(len(s)-1,math.ceil(.95*len(s))-1));return s[i]

def mean(a):return sum(a)/len(a) if a else None
def rmse(a):return math.sqrt(sum(x*x for x in a)/len(a)) if a else None
def fmt(v,n=5):
    if v is None or not math.isfinite(float(v)):return "--"
    return f"{float(v):.{n}f}"

def epoch(row):
    try:return datetime.fromisoformat(row["time_iso"]).replace(tzinfo=TZ).timestamp()
    except:return None

def qval(first,key,default=None):
    return f(first,key,default)

def clamp(v,lo,hi):
    if v is None:return None
    return max(lo,min(hi,v))

def parse_points(s):
    if not s:return []
    try:
        x=json.loads(s)
        if isinstance(x,list):return x
    except:pass
    return []

def make_xlsx(path,title,headers,rows,note):
    wb=Workbook();ws=wb.active;ws.title="Data"
    ws.append([title]+[""]*(len(headers)-1));ws.append(headers)
    for r in rows:ws.append(r)
    if len(headers)>1:ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=len(headers))
    ws["A1"].font=Font(bold=True,size=12);ws["A1"].fill=PatternFill("solid",fgColor="D9EAF7")
    for c in ws[2]:
        c.font=Font(bold=True);c.fill=PatternFill("solid",fgColor="E2F0D9");c.alignment=Alignment(horizontal="center",wrap_text=True)
    ws.freeze_panes="A3"
    for i,col in enumerate(ws.columns,1):
        vals=[len(str(c.value or "")) for c in list(col)[:100]]
        ws.column_dimensions[get_column_letter(i)].width=min(30,max(10,(max(vals) if vals else 8)+2))
    info=wb.create_sheet("Info")
    info.append(["Keterangan","Nilai"])
    info.append(["Sumber","Rekonstruksi dari RAW mission 20 Hz yang tersimpan"])
    info.append(["Tabel","Statistik kumulatif dari awal run sampai detik terkait, mengikuti konsep GUI"])
    info.append(["Catatan legacy",note])
    info.column_dimensions["A"].width=22;info.column_dimensions["B"].width=95
    wb.save(path)

def plot_lines(path,title,x,series,ylabel="Nilai"):
    fig=plt.figure(figsize=(12,5.5));ax=fig.add_subplot(111)
    any_data=False
    for label,y in series:
        vals=[float(v) if v is not None and math.isfinite(float(v)) else math.nan for v in y]
        if any(math.isfinite(v) for v in vals):
            ax.plot(x,vals,label=label);any_data=True
    ax.set_title(title);ax.set_xlabel("Waktu [s]");ax.set_ylabel(ylabel);ax.grid(True,alpha=.25)
    if any_data:ax.legend(loc="best")
    fig.tight_layout();fig.savefig(path,dpi=180);plt.close(fig)

def plot_path(path,title,pts,traj=None,goal=None):
    fig=plt.figure(figsize=(7,7));ax=fig.add_subplot(111)
    if pts:
        ax.plot([p[0] for p in pts],[p[1] for p in pts],label="Global path")
    if traj:
        ax.plot([p[0] for p in traj],[p[1] for p in traj],label="Aktual EKF global")
    if goal:
        ax.scatter([goal[0]],[goal[1]],s=70,label="Goal")
    ax.set_title(title);ax.set_xlabel("Map X [m]");ax.set_ylabel("Map Y [m]");ax.axis("equal");ax.grid(True,alpha=.25);ax.legend(loc="best")
    fig.tight_layout();fig.savefig(path,dpi=180);plt.close(fig)

def copy_overlay(src,dst,cycle=1):
    shutil.copy2(src,dst)
def main():
    ap=argparse.ArgumentParser();ap.add_argument("--raw",required=True);ap.add_argument("--package-dir",required=True);ap.add_argument("--legacy-source-dir",required=True)
    a=ap.parse_args();raw=Path(a.raw);pkg=Path(a.package_dir);legacy=Path(a.legacy_source_dir)
    rows=list(csv.DictReader(raw.open(newline="",encoding="utf-8-sig")))
    if not rows:raise SystemExit("empty raw")
    first=rows[0]
    terminal_states={"SUCCEEDED","SUCCESS","ARRIVED","FAILED","FAILURE","ABORTED","REJECTED","CANCELED","CANCELLED"}
    terminal_t=None
    for r in rows:
        st=str(r.get("goal_state.state","")).strip().upper()
        if st in terminal_states:
            terminal_t=f(r,"elapsed_s")
            if terminal_t is not None:break
    end=terminal_t if terminal_t is not None else max(f(r,"elapsed_s",0) or 0 for r in rows)
    last_sec=max(1,math.ceil(end))
    # Emulate 5 Hz GUI metric capture by selecting first RAW sample in each 0.2 s bucket.
    # Stop exactly at the Nav2 terminal result; post-terminal recorder tail is excluded.
    ticks=[];seen=set()
    for r in rows:
        t=f(r,"elapsed_s")
        if t is None or t>end:continue
        k=math.floor(t*5+1e-9)
        if k in seen:continue
        seen.add(k);ticks.append(r)
    # Static config snapshot.
    cfg={
      "lfreq":qval(first,"tuning_yaml.ekf:ekf_filter_node_odom.ros__parameters.frequency",30),
      "ltime":qval(first,"tuning_yaml.ekf:ekf_filter_node_odom.ros__parameters.sensor_timeout",.25),
      "lpredict":str(first.get("tuning_yaml.ekf:ekf_filter_node_odom.ros__parameters.predict_to_current_time","true")).lower(),
      "lqx":qval(first,"tuning_yaml.ekf:ekf_filter_node_odom.ros__parameters.process_noise_covariance.0"),
      "lqy":qval(first,"tuning_yaml.ekf:ekf_filter_node_odom.ros__parameters.process_noise_covariance.16"),
      "lqyaw":qval(first,"tuning_yaml.ekf:ekf_filter_node_odom.ros__parameters.process_noise_covariance.80"),
      "lqvx":qval(first,"tuning_yaml.ekf:ekf_filter_node_odom.ros__parameters.process_noise_covariance.96"),
      "lqwz":qval(first,"tuning_yaml.ekf:ekf_filter_node_odom.ros__parameters.process_noise_covariance.176"),
      "gfreq":qval(first,"tuning_yaml.ekf:ekf_filter_node_map.ros__parameters.frequency",30),
      "gtime":qval(first,"tuning_yaml.ekf:ekf_filter_node_map.ros__parameters.sensor_timeout",.5),
      "gpredict":str(first.get("tuning_yaml.ekf:ekf_filter_node_map.ros__parameters.predict_to_current_time","true")).lower(),
      "gqx":qval(first,"tuning_yaml.ekf:ekf_filter_node_map.ros__parameters.process_noise_covariance.0"),
      "gqy":qval(first,"tuning_yaml.ekf:ekf_filter_node_map.ros__parameters.process_noise_covariance.16"),
      "vmin":qval(first,"tuning_yaml.localization:localization_core.ros__parameters.gnss_velocity_fusion_min_variance",.0025),
      "vmax":qval(first,"tuning_yaml.localization:localization_core.ros__parameters.gnss_velocity_fusion_max_variance",.25),
      "cmin":qval(first,"tuning_yaml.localization:localization_core.ros__parameters.gnss_cog_fusion_min_variance_rad2",.00121846968),
      "cmax":qval(first,"tuning_yaml.localization:localization_core.ros__parameters.gnss_cog_fusion_max_variance_rad2",.2741556778),
      "rmin":qval(first,"tuning_yaml.nav2:planner_server.ros__parameters.GridBased.minimum_turning_radius"),
      "down":first.get("tuning_yaml.nav2:planner_server.ros__parameters.GridBased.downsample_costmap","false"),
      "factor":qval(first,"tuning_yaml.nav2:planner_server.ros__parameters.GridBased.downsampling_factor",1),
      "tol":qval(first,"tuning_yaml.nav2:planner_server.ros__parameters.GridBased.tolerance",.05),
      "cost":qval(first,"tuning_yaml.nav2:planner_server.ros__parameters.GridBased.cost_penalty",2.8),
      "xy":qval(first,"tuning_yaml.nav2:controller_server.ros__parameters.goal_checker.xy_goal_tolerance",1),
      "yawtol":qval(first,"tuning_yaml.nav2:controller_server.ros__parameters.goal_checker.yaw_goal_tolerance",math.pi),
      "stopv":qval(first,"tuning_yaml.nav2:controller_server.ros__parameters.goal_checker.trans_stopped_velocity",.08),
      "stopw":qval(first,"tuning_yaml.nav2:controller_server.ros__parameters.goal_checker.rot_stopped_velocity",.08),
    }
    # Saved legacy planner-only evidence not present in raw.
    legacy_plan_ms=515.0;legacy_clear=1.44222
    # Time-series metric samples at 5 Hz.
    S={k:[] for k in ["res_esc","res_gn","lage","lfresh","r_esc_v","r_gn_v","r_imu_yaw","r_esc_w","r_imu_w","lsx","lsy",
                       "gres","gage","gfresh","gav","hacc","r_gx","r_gy","r_gvx","r_cog","gjump","gsx","gsy",
                       "plen","pend","gdist","gspeed","gyrate","gpass"]}
    prev_sec_pos=None;sec_jump={}
    last_goal=None;goal_start=None;goal_done=None
    final_path=[];traj=[]
    for r in ticks:
        t=f(r,"elapsed_s",0);wall=epoch(r)
        # Local.
        ev=f(r,"ekf_local.v");esc=f(r,"esc_odom.v");gvx=f(r,"gnss_vel.vx")
        S["res_esc"].append((t,(esc-ev) if esc is not None and ev is not None else None))
        # Legacy fallback: raw GNSS velocity vx; fused base-frame stream was not archived.
        S["res_gn"].append((t,(gvx-ev) if gvx is not None and ev is not None else None))
        lm=f(r,"ekf_local.measurement_stamp_sec");la=max(0,wall-lm) if wall is not None and lm is not None else None
        S["lage"].append((t,la));S["lfresh"].append((t,1.0 if la is not None and la<=cfg["ltime"] else 0.0))
        S["r_esc_v"].append((t,f(r,"esc_odom.var_v")));S["r_imu_yaw"].append((t,f(r,"imu.var_yaw")));S["r_esc_w"].append((t,f(r,"esc_odom.var_w")));S["r_imu_w"].append((t,f(r,"imu.var_gz")))
        rawgv=f(r,"gnss_vel.cov_x");S["r_gn_v"].append((t,clamp(rawgv,cfg["vmin"],cfg["vmax"])))
        vx=f(r,"ekf_local.var_x");vy=f(r,"ekf_local.var_y");S["lsx"].append((t,math.sqrt(vx) if vx is not None and vx>=0 else None));S["lsy"].append((t,math.sqrt(vy) if vy is not None and vy>=0 else None))
        # Global.
        ex=f(r,"ekf_global.x");ey=f(r,"ekf_global.y");mx=f(r,"gnss_map_odom.x");my=f(r,"gnss_map_odom.y")
        gres=math.hypot(mx-ex,my-ey) if None not in (ex,ey,mx,my) else None;S["gres"].append((t,gres))
        gm=f(r,"ekf_global.measurement_stamp_sec");ga=max(0,wall-gm) if wall is not None and gm is not None else None
        S["gage"].append((t,ga));S["gfresh"].append((t,1.0 if ga is not None and ga<=cfg["gtime"] else 0.0))
        fm=f(r,"gnss_fix.measurement_stamp_sec");fa=max(0,wall-fm) if wall is not None and fm is not None else None
        fixok=b(r,"gnss_quality.gnss_fix_ok",True);S["gav"].append((t,1.0 if fa is not None and fa<=1.5 and fixok else 0.0))
        S["hacc"].append((t,f(r,"gnss_quality.hacc_m")));S["r_gx"].append((t,f(r,"gnss_map_odom.var_x")));S["r_gy"].append((t,f(r,"gnss_map_odom.var_y")))
        S["r_gvx"].append((t,clamp(f(r,"gnss_vel.cov_x"),cfg["vmin"],cfg["vmax"])))
        ca=f(r,"gnss_quality.course_accuracy_rad");S["r_cog"].append((t,clamp(ca*ca if ca is not None else cfg["cmax"],cfg["cmin"],cfg["cmax"])))
        sec=int(math.floor(t+1e-9))
        if ex is not None and ey is not None:
            if sec not in sec_jump:
                j=0.0 if prev_sec_pos is None else math.hypot(ex-prev_sec_pos[0],ey-prev_sec_pos[1]);sec_jump[sec]=j;prev_sec_pos=(ex,ey)
            S["gjump"].append((t,sec_jump[sec]));traj.append((ex,ey))
        vx=f(r,"ekf_global.var_x");vy=f(r,"ekf_global.var_y");S["gsx"].append((t,math.sqrt(vx) if vx is not None and vx>=0 else None));S["gsy"].append((t,math.sqrt(vy) if vy is not None and vy>=0 else None))
        plen=f(r,"nav_path.length_m");S["plen"].append((t,plen))
        pts=parse_points(r.get("nav_path.points",""));gx=f(r,"goal_pose.x");gy=f(r,"goal_pose.y")
        if pts:final_path=pts
        pend=math.hypot(float(pts[-1][0])-gx,float(pts[-1][1])-gy) if pts and gx is not None and gy is not None else None;S["pend"].append((t,pend))
        if gx is not None and gy is not None and ex is not None and ey is not None:
            gd=math.hypot(ex-gx,ey-gy);S["gdist"].append((t,gd))
        sp=abs(f(r,"esc_drive_actual",f(r,"ekf_global.v",0)) or 0);yw=abs(f(r,"esc_yaw_rate",f(r,"ekf_global.w",0)) or 0);S["gspeed"].append((t,sp));S["gyrate"].append((t,yw))
        gyaw=f(r,"goal_pose.yaw");eyaw=f(r,"ekf_global.yaw");yawerr=abs(((eyaw-gyaw+math.pi)%(2*math.pi))-math.pi) if eyaw is not None and gyaw is not None else None
        gd=S["gdist"][-1][1] if S["gdist"] else None
        passed=1.0 if gd is not None and gd<=cfg["xy"] and yawerr is not None and yawerr<=cfg["yawtol"] and sp<=cfg["stopv"] and yw<=cfg["stopw"] else 0.0
        S["gpass"].append((t,passed))
        st=str(r.get("goal_state.state","")).strip().upper()
        if st and st!=last_goal:
            if st not in {"IDLE","UNKNOWN",""} and goal_start is None:goal_start=t
            if st in {"SUCCEEDED","SUCCESS","ARRIVED","FAILED","ABORTED","CANCELED","CANCELLED","REJECTED"}:goal_done=t
            last_goal=st
    def vals(key,sec):return [v for t,v in S[key] if t<=sec and v is not None and math.isfinite(float(v))]
    def last(key,sec):
        a=vals(key,sec);return a[-1] if a else None
    def longest_zero(sec):
        pts=[(t,v) for t,v in S["gav"] if t<=sec and v is not None];best=0;start=None;lastt=None
        for t,v in pts:
            if v<.5:
                if start is None:start=t
                lastt=t
            elif start is not None:
                best=max(best,(lastt-start)+.2);start=lastt=None
        if start is not None:best=max(best,(lastt-start)+.2)
        return best
    variation="run-01"
    table_rows={sid:[] for sid in TABLES}
    for sec in range(1,last_sec+1):
        re_=vals("res_esc",sec);rg=vals("res_gn",sec);la=vals("lage",sec);lf=vals("lfresh",sec)
        gr=vals("gres",sec);ga=vals("gage",sec);gf=vals("gfresh",sec);av=vals("gav",sec);j=vals("gjump",sec)
        pathlen=last("plen",sec);pend=last("pend",sec);planner="PASS" if pathlen is not None and pathlen>0 else "WAIT"
        planms=legacy_plan_ms if planner=="PASS" else None;clear=legacy_clear if planner=="PASS" else None
        gd=vals("gdist",sec);gs=vals("gspeed",sec);gw=vals("gyrate",sec);gp=vals("gpass",sec)
        state="WAIT"
        raw_at=[r for r in rows if (f(r,"elapsed_s",-1) or -1)<=sec]
        if raw_at: state=str(raw_at[-1].get("goal_state.state","") or "WAIT").upper()
        tt=(min(sec,goal_done)-goal_start) if goal_start is not None and goal_done is not None and sec>=goal_start else (sec-goal_start if goal_start is not None else None)
        table_rows["4.2.1"].append([sec,variation,fmt(cfg["lfreq"],6),fmt(rmse(re_)),fmt(rmse(rg)),fmt(mean(la)),fmt(p95(la)),fmt(mean(lf)*100 if lf else None,1)])
        table_rows["4.2.2"].append([sec,variation,fmt(cfg["ltime"],6),fmt(rmse(re_)),fmt(rmse(rg)),fmt(mean(la)),fmt(p95(la)),fmt(mean(lf)*100 if lf else None,1)])
        table_rows["4.2.3"].append([sec,variation,fmt(mean(vals("r_esc_v",sec))),fmt(mean(vals("r_gn_v",sec))),fmt(mean(vals("r_imu_yaw",sec))),fmt(mean(vals("r_esc_w",sec))),fmt(mean(vals("r_imu_w",sec))),fmt(rmse(re_)),fmt(rmse(rg))])
        table_rows["4.2.4"].append([sec,variation,fmt(cfg["lqx"],6),fmt(cfg["lqy"],6),fmt(cfg["lqyaw"],6),fmt(cfg["lqvx"],6),fmt(cfg["lqwz"],6),fmt(mean(vals("lsx",sec))),fmt(mean(vals("lsy",sec))),fmt(rmse(re_)),fmt(rmse(rg))])
        table_rows["4.2.5"].append([sec,variation,cfg["lpredict"],fmt(rmse(re_)),fmt(rmse(rg)),fmt(mean(la)),fmt(p95(la)),fmt(mean(lf)*100 if lf else None,1)])
        table_rows["4.3.1"].append([sec,variation,fmt(cfg["gfreq"],6),fmt(rmse(gr)),fmt(mean(ga)),fmt(p95(ga)),fmt(mean(gf)*100 if gf else None,1),fmt(mean(av)*100 if av else None,1)])
        table_rows["4.3.2"].append([sec,variation,fmt(cfg["gtime"],6),fmt(mean(av)*100 if av else None,1),fmt(longest_zero(sec),1),fmt(mean(gf)*100 if gf else None,1),fmt(rmse(gr)),fmt(mean(ga))])
        table_rows["4.3.3"].append([sec,variation,fmt(mean(vals("hacc",sec)),3),fmt(mean(vals("r_gx",sec))),fmt(mean(vals("r_gy",sec))),fmt(mean(vals("r_gvx",sec))),fmt(mean(vals("r_cog",sec))),fmt(rmse(gr)),fmt(max(j) if j else None)])
        table_rows["4.3.4"].append([sec,variation,fmt(cfg["gqx"],6),fmt(cfg["gqy"],6),fmt(mean(vals("gsx",sec))),fmt(mean(vals("gsy",sec))),fmt(rmse(gr)),fmt(max(j) if j else None)])
        table_rows["4.3.5"].append([sec,variation,cfg["gpredict"],fmt(rmse(gr)),fmt(max(j) if j else None),fmt(mean(ga)),fmt(p95(ga)),fmt(mean(gf)*100 if gf else None,1)])
        table_rows["4.4.1"].append([sec,variation,fmt(cfg["rmin"],6),fmt(planms),fmt(pathlen),fmt(pend),fmt(clear),planner])
        table_rows["4.4.2"].append([sec,variation,cfg["down"],fmt(cfg["factor"],6),fmt(planms),fmt(pathlen),fmt(clear),planner])
        table_rows["4.4.3"].append([sec,variation,fmt(cfg["tol"],6),fmt(pend),fmt(planms),planner,fmt(pathlen)])
        table_rows["4.4.4"].append([sec,variation,fmt(cfg["cost"],6),fmt(clear),fmt(pathlen),fmt(planms),planner])
        table_rows["4.4.5"].append([sec,variation,fmt(cfg["xy"],6),fmt(cfg["yawtol"],6),fmt(cfg["stopv"],6),fmt(cfg["stopw"],6),fmt(gd[-1] if gd else None),fmt(min(gd) if gd else None),fmt(gs[-1] if gs else None),fmt(gw[-1] if gw else None),("PASS" if gp and gp[-1]>=.5 else "WAIT"),state,fmt(tt,3)])
    # Rebuild report tree only; preserve raw/internal evidence elsewhere.
    report=pkg/"report"
    if report.exists():shutil.rmtree(report)
    note=("Legacy success mission reconstructed from preserved RAW 20 Hz. GUI 1 Hz table capture was incomplete in the original run. "
          "GNSS local residual uses archived GNSS velocity vx fallback because gnss_base_vel_fusion stream was not archived. "
          "Planner latency/min-clearance use the saved final evidence (515 ms / 1.44222 m) once a valid path exists.")
    # For graph source, use reconstructed instantaneous/cumulative series.
    secx=list(range(1,last_sec+1))
    def colnums(sid,col):
        idx=TABLES[sid][1].index(col);out=[]
        for r in table_rows[sid]:
            try:out.append(float(r[idx]))
            except:out.append(None)
        return out
    for sid,(label,headers) in TABLES.items():
        folder=report/(".".join(sid.split(".")[:2]))/sid;folder.mkdir(parents=True,exist_ok=True)
        xp=folder/f"Tabel_{label}.xlsx";make_xlsx(xp,f"Tabel {sid} - {label}",headers,table_rows[sid],note)
        # Primary report graph matching the subchapter's GUI content.
        gp=folder/f"Grafik_{label}.png"
        if sid=="4.2.1":plot_lines(gp,"Pengaruh frequency terhadap respons EKF lokal",secx,[("RMSE Residual vx ESC",colnums(sid,"RMSE Residual vx ESC")),("RMSE Residual vx GNSS",colnums(sid,"RMSE Residual vx GNSS")),("Mean EKF age",colnums(sid,"Mean EKF age"))],"Residual / state age")
        elif sid=="4.2.2":plot_lines(gp,"Pengaruh sensor_timeout EKF lokal",secx,[("Mean EKF age",colnums(sid,"Mean EKF age")),("EKF continuity [%]",colnums(sid,"EKF continuity [%]"))],"State age / continuity")
        elif sid=="4.2.3":plot_lines(gp,"Sensitivitas measurement covariance EKF lokal",secx,[("R ESC vx",colnums(sid,"R ESC vx")),("R GNSS vx",colnums(sid,"R GNSS vx")),("R IMU wz",colnums(sid,"R IMU wz"))],"Measurement variance")
        elif sid=="4.2.4":plot_lines(gp,"Pengaruh process noise EKF lokal",secx,[("Mean sigma X",colnums(sid,"Mean sigma X")),("Mean sigma Y",colnums(sid,"Mean sigma Y")),("RMSE Residual vx ESC",colnums(sid,"RMSE Residual vx ESC"))],"Uncertainty / residual")
        elif sid=="4.2.5":plot_lines(gp,"Pengaruh predict_to_current_time EKF lokal",secx,[("Mean EKF age",colnums(sid,"Mean EKF age")),("RMSE Residual vx ESC",colnums(sid,"RMSE Residual vx ESC"))],"State age / residual")
        elif sid=="4.3.1":plot_lines(gp,"Pengaruh frequency EKF global",secx,[("RMSE Residual XY GNSS [m]",colnums(sid,"RMSE Residual XY GNSS [m]")),("Mean EKF age",colnums(sid,"Mean EKF age"))],"Residual posisi / state age")
        elif sid=="4.3.2":plot_lines(gp,"Pengaruh sensor_timeout EKF global",secx,[("GNSS availability [%]",colnums(sid,"GNSS availability [%]")),("EKF global continuity [%]",colnums(sid,"EKF global continuity [%]")),("RMSE Residual XY GNSS [m]",colnums(sid,"RMSE Residual XY GNSS [m]"))],"Availability / residual")
        elif sid=="4.3.3":plot_lines(gp,"Sensitivitas measurement covariance GNSS",secx,[("R GNSS X",colnums(sid,"R GNSS X")),("R GNSS Y",colnums(sid,"R GNSS Y")),("RMSE Residual XY GNSS [m]",colnums(sid,"RMSE Residual XY GNSS [m]"))],"Position variance / residual")
        elif sid=="4.3.4":plot_lines(gp,"Pengaruh process noise posisi EKF global",secx,[("Mean sigma X",colnums(sid,"Mean sigma X")),("Mean sigma Y",colnums(sid,"Mean sigma Y")),("RMSE Residual XY GNSS [m]",colnums(sid,"RMSE Residual XY GNSS [m]"))],"Uncertainty / residual [m]")
        elif sid=="4.3.5":plot_lines(gp,"Pengaruh predict_to_current_time EKF global",secx,[("RMSE Residual XY GNSS [m]",colnums(sid,"RMSE Residual XY GNSS [m]")),("Max position jump [m]",colnums(sid,"Max position jump [m]")),("Mean EKF age",colnums(sid,"Mean EKF age"))],"Residual / jump / state age")
        elif sid=="4.4.1":
            gx=f(rows[-1],"goal_pose.x");gy=f(rows[-1],"goal_pose.y");plot_path(gp,"Pengaruh minimum turning radius",final_path,traj,(gx,gy) if gx is not None and gy is not None else None)
        elif sid=="4.4.2":plot_lines(gp,"Trade-off downsampling factor",secx,[("Planning time",colnums(sid,"Planning time")),("Minimum clearance [m]",colnums(sid,"Minimum clearance [m]"))],"Planning time / clearance")
        elif sid=="4.4.3":plot_lines(gp,"Pengaruh planner tolerance terhadap endpoint error",secx,[("Path endpoint error",colnums(sid,"Path endpoint error")),("Planning time",colnums(sid,"Planning time"))],"Endpoint error / planning time")
        elif sid=="4.4.4":plot_lines(gp,"Pengaruh cost_penalty",secx,[("Minimum clearance [m]",colnums(sid,"Minimum clearance [m]")),("Path length",colnums(sid,"Path length")),("Planning time",colnums(sid,"Planning time"))],"Clearance / path / planning")
        elif sid=="4.4.5":
            plot_lines(gp,"Jarak kendaraan terhadap radius Goal Checker",secx,[("Distance to goal",colnums(sid,"Final distance to goal")),("XY tolerance",[cfg["xy"]]*len(secx))],"Jarak [m]")
            plot_lines(folder/f"Grafik_2_{label}.png","Kecepatan translasi terhadap stopped threshold",secx,[("Speed abs",colnums(sid,"Final speed abs")),("Stop speed threshold",[cfg["stopv"]]*len(secx))],"Kecepatan [m/s]")
            plot_lines(folder/f"Grafik_3_{label}.png","Yaw-rate terhadap stopped threshold",secx,[("Yaw rate abs",colnums(sid,"Final yaw rate abs")),("Stop yaw threshold",[cfg["stopw"]]*len(secx))],"Yaw rate [rad/s]")
            gx=f(rows[-1],"goal_pose.x");gy=f(rows[-1],"goal_pose.y");plot_path(folder/f"Grafik_4_{label}.png","Trajectory aktual pendekatan goal",final_path,traj,(gx,gy) if gx is not None and gy is not None else None)
        copy_overlay(gp,folder/f"Overlay_{label}.png")
    # Zip only 4.2 / 4.3 / 4.4.
    zpath=pkg/"BAB4_Report_Cycle_01.zip"
    zpath.unlink(missing_ok=True)
    with zipfile.ZipFile(zpath,"w",zipfile.ZIP_DEFLATED,compresslevel=6) as z:
        for p in sorted(report.rglob("*")):
            if p.is_file():z.write(p,p.relative_to(report))
    meta_path=pkg/"package.json";meta=json.load(open(meta_path));meta["zip_path"]=str(zpath);meta["zip_name"]=zpath.name;meta["bytes"]=zpath.stat().st_size;meta["legacy_reconstructed_full_seconds"]=True;meta["reconstructed_seconds"]=last_sec
    with open(meta_path,"w") as f_:json.dump(meta,f_,ensure_ascii=False,indent=2);f_.write("\n")
    print(json.dumps({"seconds":last_sec,"zip":str(zpath),"bytes":zpath.stat().st_size},ensure_ascii=False))

if __name__=="__main__":main()
